"""dict.json → 통계용어사전_영한_한영.xlsx

사용법: python scripts/export_xlsx.py [--include-deleted] [--out 경로]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).parent))
from dictcore import en_key  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
STATUS_KO = {"original": "원천", "added": "추가", "updated": "수정", "deleted": "삭제됨"}
FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F4B5A")


def en_sort(s):
    return re.sub(r"[^0-9a-z ]", "", en_key(s))


def add_sheet(wb, title, header, rows, widths, table_name):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    if rows:
        t = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(header))}{len(rows) + 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dict", default=ROOT / "web" / "data" / "dict.json", type=Path)
    ap.add_argument("--edits", default=ROOT / "data" / "edits.json", type=Path)
    ap.add_argument("--out", default=ROOT / "dist" / "통계용어사전_영한_한영.xlsx", type=Path)
    ap.add_argument("--include-deleted", action="store_true")
    a = ap.parse_args()

    doc = json.loads(a.dict.read_text(encoding="utf-8"))
    meta = doc["meta"]
    entries = doc["entries"]
    active = [e for e in entries if e.get("status") != "deleted"]
    st = lambda e: STATUS_KO[e.get("status", "original")]  # noqa: E731
    src = lambda e: e.get("src_row") or ""  # noqa: E731

    # 영한 ---------------------------------------------------------------
    groups = {}
    for e in active:
        k = (en_key(e["en"]), e.get("sense") or 0)
        groups.setdefault(k, []).append(e)
    en_rows = []
    for (_, sense), es in sorted(groups.items(), key=lambda kv: (en_sort(kv[1][0]["en"]), kv[0][1])):
        forms = list(dict.fromkeys(f for e in es for f in e["ko_forms"]))
        en_rows.append([
            es[0]["en"],
            ", ".join(dict.fromkeys(v for e in es for v in e.get("en_variants", []))),
            ", ".join(dict.fromkeys(v for e in es for v in e.get("abbr", []))),
            sense or "",
            "; ".join(dict.fromkeys(e["ko_primary"] for e in es)),
            "; ".join(forms),
            ", ".join(dict.fromkeys(st(e) for e in es)),
            ", ".join(str(src(e)) for e in es if src(e)),
            ", ".join(e["id"] for e in es),
        ])

    # 한영 ---------------------------------------------------------------
    ko_map = {}
    for e in active:
        for f in e["ko_forms"]:
            ko_map.setdefault(f, []).append(e)
    ko_rows = []
    for f in sorted(ko_map, key=lambda s: (not re.match(r"[\uac00-\ud7a3]", s), s)):
        es = ko_map[f]
        ens = []
        for e in es:
            label = e["en"] + (f" ({e['sense']})" if e.get("sense") else "")
            if e.get("abbr"):
                label += f" [{', '.join(e['abbr'])}]"
            ens.append(label)
        syn = dict.fromkeys(x for e in es for x in e["ko_forms"] if x != f)
        ko_rows.append([
            f,
            "예" if any(e["ko_primary"] == f for e in es) else "",
            "; ".join(dict.fromkeys(ens)),
            ", ".join(syn),
            "; ".join(dict.fromkeys(n for e in es for n in e.get("notes", []))),
            ", ".join(dict.fromkeys(st(e) for e in es)),
            ", ".join(str(src(e)) for e in es if src(e)),
            ", ".join(e["id"] for e in es),
        ])

    # 검토필요 -------------------------------------------------------------
    review = []
    for e in entries:
        fl = [x for x in e.get("flags", []) if x in ("paren_unreviewed", "expansion_capped", "unbalanced", "en_fixed")]
        if fl:
            review.append([e["id"], e["ko_raw"], e["en_raw"], "; ".join(e["ko_forms"]), ", ".join(fl)])

    # 변경이력 -------------------------------------------------------------
    edits = []
    if a.edits.exists():
        data = json.loads(a.edits.read_text(encoding="utf-8") or "[]")
        edits = data.get("edits", data) if isinstance(data, dict) else data
    OPS = {"add": "추가", "update": "수정", "delete": "삭제", "restore": "복원"}
    hist = [[x.get("edit_id"), x.get("created_at"), OPS.get(x.get("op"), x.get("op")), x.get("target_id"),
             (x.get("before") or {}).get("ko_raw", ""), (x.get("before") or {}).get("en_raw", ""),
             (x.get("after") or {}).get("ko_raw", ""), (x.get("after") or {}).get("en_raw", ""),
             x.get("reason", ""), x.get("editor", ""), x.get("status", "")]
            for x in sorted(edits, key=lambda x: (x.get("created_at", ""), x.get("edit_id", "")))]

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "영한", ["영문", "영문 이형", "약어", "의미번호", "한글 역어(대표)", "한글 역어(전체)", "상태", "원천 행", "id"],
              en_rows, [40, 22, 14, 9, 26, 50, 10, 12, 16], "EnKo")
    add_sheet(wb, "한영", ["한글 표기", "대표 표기", "영문", "동의 역어", "주석(한자 등)", "상태", "원천 행", "id"],
              ko_rows, [30, 9, 44, 40, 14, 10, 12, 16], "KoEn")
    add_sheet(wb, "검토필요", ["id", "한글 원문", "영문 원문", "전개 결과", "사유"],
              review, [10, 34, 40, 50, 30], "Review")
    add_sheet(wb, "변경이력", ["편집 id", "시각", "연산", "대상 id", "이전 한글", "이전 영문", "이후 한글", "이후 영문",
                           "사유", "편집자", "상태"], hist, [16, 24, 8, 10, 24, 24, 24, 24, 24, 12, 10], "History")
    if a.include_deleted:
        dele = [[e["id"], e["ko_raw"], e["en_raw"], e.get("delete_reason") or "", e.get("src_row") or ""]
                for e in entries if e.get("status") == "deleted"]
        add_sheet(wb, "삭제됨", ["id", "한글 원문", "영문 원문", "삭제 사유", "원천 행"], dele,
                  [10, 34, 40, 30, 10], "Deleted")

    info = wb.create_sheet("정보", 0)
    info["A1"] = "통계용어사전 (영한·한영)"
    info["A1"].font = Font(name=FONT, bold=True, size=14)
    rows = [
        ("출처", "한국통계학회 통계용어 (https://kss.or.kr/homepage/custom/statistics)"),
        ("원천 파일", meta["source_file"]),
        ("원천 SHA-256", meta["source_sha256"]),
        ("원천 행 / 유효 행", f"{meta['source_rows']} / {meta['rows']}"),
        ("편집 적용 / 충돌", f"{meta['edits_applied']} / {meta['conflicts']}"),
        ("안내", "'상태'가 추가·수정인 항목은 학회 원본과 내용이 다릅니다."),
    ]
    for i, (k, v) in enumerate(rows, 3):
        info.cell(i, 1, k).font = Font(name=FONT, bold=True)
        info.cell(i, 2, v).font = Font(name=FONT)
    r0 = len(rows) + 4
    counts = [
        ("영한 표제어 수", "=COUNTA(영한!A:A)-1"),
        ("한영 표제어 수", "=COUNTA(한영!A:A)-1"),
        ("추가된 항목(영한 기준)", '=COUNTIF(영한!G:G,"*추가*")'),
        ("수정된 항목(영한 기준)", '=COUNTIF(영한!G:G,"*수정*")'),
        ("검토 필요 항목", "=COUNTA(검토필요!A:A)-1"),
    ]
    for i, (k, f) in enumerate(counts, r0):
        info.cell(i, 1, k).font = Font(name=FONT, bold=True)
        info.cell(i, 2, f).font = Font(name=FONT)
        info.cell(i, 2).alignment = Alignment(horizontal="left")
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 80

    a.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    print(f"영한 {len(en_rows)} · 한영 {len(ko_rows)} · 검토필요 {len(review)} · 변경이력 {len(hist)} → {a.out}")


if __name__ == "__main__":
    main()
