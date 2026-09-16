"""편집요청.xlsx → data/edits.json (일괄 편집, PRD E11)

  python scripts/import_edits.py --template            # 빈 템플릿 생성
  python scripts/import_edits.py 편집요청.xlsx          # 검증 후 edits.json에 추가
  python scripts/import_edits.py 편집요청.xlsx --dry-run --draft
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from dictcore import ExpandError, apply_edits, expand_ko, normalize_en  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
KST = timezone(timedelta(hours=9))
OPS = {"추가": "add", "수정": "update", "삭제": "delete", "복원": "restore"}
PT = {"생략": "optional", "대체": "alternative", "주석": "annotation",
      "optional": "optional", "alternative": "alternative", "annotation": "annotation"}
HEADER = ["연산", "대상 id", "한글 원문", "영문 원문", "소괄호 유형", "뜻풀이", "출처 메모", "사유", "편집자"]


def make_template(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "편집요청"
    ws.append(HEADER)
    fill = PatternFill("solid", start_color="FFF2CC")
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True)
        c.fill = fill
    for col, w in zip("ABCDEFGHI", [8, 10, 30, 36, 16, 40, 16, 24, 10]):
        ws.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1='"추가,수정,삭제,복원"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("A2:A1000")
    ws.freeze_panes = "A2"

    ex = wb.create_sheet("작성 예시")
    ex.append(HEADER)
    for r in [
        ["추가", "", "등각예측, 컨포멀 예측", "conformal prediction", "", "예측구간을 분포 가정 없이 구성하는 방법", "", "신규 용어", "홍길동"],
        ["수정", "r0003", "가능도, 우도, 라이클리후드", "", "", "", "", "외래어 표기 추가", "홍길동"],
        ["수정", "r0012", "", "", "대체", "", "", "소괄호 유형 지정", "홍길동"],
        ["삭제", "r0036", "", "", "", "", "", "중복 항목", "홍길동"],
    ]:
        ex.append(r)
    for c in ex[1]:
        c.font = Font(name="Arial", bold=True)

    g = wb.create_sheet("안내")
    for line in [
        "'편집요청' 시트의 노란 머리글 아래에 한 줄에 편집 1건씩 적습니다. '작성 예시' 시트는 읽지 않습니다.",
        "연산: 추가 / 수정 / 삭제 / 복원",
        "대상 id: 수정·삭제·복원은 필수(r0003, u0001 형식). 추가는 비워 두면 새 id가 자동 발급됩니다.",
        "한글·영문 원문: 추가는 필수. 수정에서 빈칸은 '바꾸지 않음'입니다.",
        "한글 원문 표기법: 쉼표(동의어), [대체어], (소괄호). 예: 가능도[우도]비 검정[검증]",
        "소괄호 유형: 소괄호 순서대로 생략/대체/주석을 세미콜론으로 구분 (예: 생략;대체)",
        "사유: 변경이력에 남습니다. 삭제는 사유를 꼭 적어 주세요.",
    ]:
        g.append([line])
    g.column_dimensions["A"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"템플릿 생성: {path}")


def load_json(path: Path, default):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def text(v):
    return re.sub(r"\s+", " ", str(v).replace("\u00a0", " ")).strip() if v is not None else ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", type=Path)
    ap.add_argument("--template", action="store_true")
    ap.add_argument("--dict", default=ROOT / "web" / "data" / "dict.json", type=Path)
    ap.add_argument("--edits", default=ROOT / "data" / "edits.json", type=Path)
    ap.add_argument("--draft", action="store_true", help="status=draft 로 추가 (빌드 미반영)")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    if a.template:
        make_template(a.xlsx or ROOT / "data" / "편집요청.xlsx")
        return
    if not a.xlsx:
        ap.error("편집요청 엑셀 경로를 지정하세요 (또는 --template)")

    doc = load_json(a.dict, {"entries": []})
    entries = {e["id"]: e for e in doc["entries"]}
    store = load_json(a.edits, {"version": 1, "edits": []})
    existing = store["edits"] if isinstance(store, dict) else store
    # 이미 빌드되지 않은 기존 편집(draft 포함)까지 반영한 현재 상태에서 검증
    state, _, _ = apply_edits(copy.deepcopy(entries), copy.deepcopy(existing), only_approved=False)

    used = [int(m.group(1)) for i in list(state) + [e.get("target_id", "") for e in existing]
            if (m := re.fullmatch(r"u(\d+)", i or ""))]
    next_u = max(used, default=0) + 1
    today = datetime.now(KST)
    seq_base = sum(1 for e in existing if (e.get("edit_id") or "").startswith(f"e{today:%Y%m%d}-"))

    ws = load_workbook(a.xlsx, data_only=True)["편집요청"]
    new, errors = [], []
    for rn, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        vals = dict(zip(HEADER, [text(v) for v in row[: len(HEADER)]]))
        if not any(vals.values()):
            continue
        op = OPS.get(vals["연산"])
        tid = vals["대상 id"]
        err = lambda m: errors.append(f"{rn}행: {m}")  # noqa: E731
        if not op:
            err(f"연산을 알 수 없음 '{vals['연산']}'")
            continue
        ptypes = None
        if vals["소괄호 유형"]:
            try:
                ptypes = [PT[x.strip()] for x in vals["소괄호 유형"].split(";") if x.strip()]
            except KeyError as ex:
                err(f"소괄호 유형 오류 {ex}")
                continue
        after = {}
        if op == "add":
            if not tid:
                tid = f"u{next_u:04d}"
                next_u += 1
            if tid in state:
                err(f"id {tid} 는 이미 존재")
                continue
            after = {"ko_raw": vals["한글 원문"], "en_raw": vals["영문 원문"]}
        else:
            cur = state.get(tid)
            if not cur:
                err(f"대상 id '{tid}' 없음")
                continue
            if op == "update":
                for k, col in (("ko_raw", "한글 원문"), ("en_raw", "영문 원문"),
                               ("definition", "뜻풀이"), ("source_note", "출처 메모")):
                    if vals[col]:
                        after[k] = vals[col]
                if ptypes is not None:
                    after["paren_types"] = ptypes
                if not after:
                    err("수정할 내용이 없음")
                    continue
        if op in ("add",):
            for k, col in (("definition", "뜻풀이"), ("source_note", "출처 메모")):
                if vals[col]:
                    after[k] = vals[col]
            if ptypes is not None:
                after["paren_types"] = ptypes
        if op in ("add", "update"):
            probe = state.get(tid, {})
            try:
                expand_ko(after.get("ko_raw", probe.get("ko_raw", "")),
                          after.get("paren_types", probe.get("paren_types")), strict=True)
                normalize_en(after.get("en_raw", probe.get("en_raw", "")))
            except ExpandError as ex:
                err(str(ex))
                continue
        if op == "delete" and not vals["사유"]:
            err("삭제 사유가 필요함")
            continue
        seq_base += 1
        cur = state.get(tid) or {}
        edit = {
            "edit_id": f"e{today:%Y%m%d}-{seq_base:04d}",
            "op": op,
            "target_id": tid,
            "before": {"ko_raw": cur["ko_raw"], "en_raw": cur["en_raw"]} if cur else {},
            "after": after,
            "reason": vals["사유"],
            "editor": vals["편집자"],
            "created_at": (today + timedelta(milliseconds=rn)).isoformat(timespec="milliseconds"),
            "status": "draft" if a.draft else "approved",
        }
        state, conflicts, _ = apply_edits(state, [edit], only_approved=False)
        if conflicts:
            err(conflicts[0]["reason"])
            continue
        new.append(edit)

    for e in errors:
        print("오류 -", e)
    print(f"유효 편집 {len(new)}건, 오류 {len(errors)}건")
    if errors:
        sys.exit("오류가 있어 edits.json을 바꾸지 않았습니다. 엑셀을 고친 뒤 다시 실행하세요.")
    if a.dry_run:
        print(json.dumps(new, ensure_ascii=False, indent=1)[:2000])
        return
    out = {"version": 1, "edits": existing + new}
    a.edits.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"→ {a.edits} (총 {len(out['edits'])}건). 이제 build_dict.py 를 실행하세요.")


if __name__ == "__main__":
    main()
